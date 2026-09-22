"""Blue Ocean Finder — Word(.docx) 보고서 생성기.

`generate_blue_ocean_report(data, filename=...)` 하나만 외부에 노출한다.
`data` 딕셔너리에 없는 값은 절대 추측해서 채우지 않고 "확인되지 않음"으로 표시한다.

허용 키:
    country, blue_ocean_score, market_opportunity_score, penetration_opportunity_score,
    growth_1y, cagr_3y, korea_market_share, global_korea_share, export_gap,
    tariff_rate, top3_concentration, item_name, hs_code, project_name, date

주의: "종합 순위"(rank)는 위 스키마에 없는 선택 필드다. 호출 측이 `data["rank"]`를
추가로 넘기면 표시하고, 없으면 항상 "확인되지 않음"으로 남는다 — 지어내지 않기 위함.
"""
from __future__ import annotations

from datetime import date

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Inches, Pt, RGBColor

# ─────────────────────────────────────────────────────────────────────────
# 색상 팔레트
# ─────────────────────────────────────────────────────────────────────────
NAVY_DEEP = "0F2240"       # 딥 네이비 (주 헤더, 타이틀)
NAVY_PRIMARY = "1B365D"    # 프라이머리 네이비 (표 헤더, 주요 수치)
BLUE_ACCENT = "2563AF"     # 액센트 블루 (라인, 배지 라벨)
BG_LIGHT = "F8FAFC"
BG_LIGHT2 = "EEF4FA"
TEXT_BODY = "212936"
TEXT_SUB = "64748B"
STATUS_GOOD_TEXT, STATUS_GOOD_BG = "166534", "E8F3EC"
STATUS_WARN_TEXT, STATUS_WARN_BG = "B45309", "FBF3DF"
# 팔레트에 "상(위험)" 색이 없어, 신호등 표현을 위해 같은 톤으로 자체 확장한 색.
STATUS_RISK_TEXT, STATUS_RISK_BG = "B91C1C", "FDECEC"
WHITE = "FFFFFF"
GRAY_BORDER = "D9E2EC"

SEVERITY_COLOR = {
    "하": (STATUS_GOOD_BG, STATUS_GOOD_TEXT),
    "중": (STATUS_WARN_BG, STATUS_WARN_TEXT),
    "상": (STATUS_RISK_BG, STATUS_RISK_TEXT),
}

FONT_NAME = "맑은 고딕"
CONTENT_WIDTH = Inches(6.7)


# ─────────────────────────────────────────────────────────────────────────
# XML 저수준 헬퍼 (셀 배경 · 여백 · 테두리를 parse_xml로 정밀 제어)
# ─────────────────────────────────────────────────────────────────────────
def _c(hex_color: str) -> str:
    return hex_color.lstrip("#")


def set_cell_shading(cell, hex_color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{_c(hex_color)}"/>'))


def set_cell_margins(cell, top=60, bottom=60, left=100, right=100) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    xml = (
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f"</w:tcMar>"
    )
    tcPr.append(parse_xml(xml))


def set_cell_borders(cell, **edges) -> None:
    """edges: top/bottom/left/right -> {"sz": int, "val": str, "color": str}"""
    tcPr = cell._tc.get_or_add_tcPr()
    parts = []
    for edge, spec in edges.items():
        sz = spec.get("sz", 4)
        val = spec.get("val", "single")
        color = _c(spec.get("color", "000000"))
        parts.append(f'<w:{edge} w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>')
    tcPr.append(parse_xml(f'<w:tcBorders {nsdecls("w")}>' + "".join(parts) + "</w:tcBorders>"))


def set_table_borders(table, **edges) -> None:
    tblPr = table._tbl.get_or_add_tblPr()
    parts = []
    for edge, spec in edges.items():
        sz = spec.get("sz", 4)
        val = spec.get("val", "single")
        color = _c(spec.get("color", "000000"))
        parts.append(f'<w:{edge} w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>')
    tblPr.append(parse_xml(f'<w:tblBorders {nsdecls("w")}>' + "".join(parts) + "</w:tblBorders>"))


def _apply_font(run, name: str = FONT_NAME) -> None:
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)


def _tight(paragraph, before=2, after=2, line=1.12) -> None:
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def _run(paragraph, text, size=9.5, bold=False, color=TEXT_BODY, italic=False):
    r = paragraph.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = RGBColor.from_string(_c(color))
    _apply_font(r)
    return r


def _set_col_widths(table, widths) -> None:
    table.autofit = False
    for row in table.rows:
        for idx, w in enumerate(widths):
            row.cells[idx].width = w


# ─────────────────────────────────────────────────────────────────────────
# 값 포맷 (없는 값은 절대 지어내지 않고 "확인되지 않음")
# ─────────────────────────────────────────────────────────────────────────
def fmt(value, unit="", decimals=1, signed=False, na="확인되지 않음") -> str:
    if value is None:
        return na
    try:
        v = float(value)
    except (TypeError, ValueError):
        return na
    sign = "+" if (signed and v > 0) else ""
    return f"{sign}{v:.{decimals}f}{unit}"


def safe(value, na="확인되지 않음"):
    return value if value not in (None, "") else na


# ─────────────────────────────────────────────────────────────────────────
# 조립 부품
# ─────────────────────────────────────────────────────────────────────────
def add_divider(doc, width=Inches(2.5), color=BLUE_ACCENT):
    table = doc.add_table(rows=1, cols=1)
    _set_col_widths(table, [width])
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, color)
    set_cell_margins(cell, top=8, bottom=8, left=0, right=0)
    p = cell.paragraphs[0]
    _tight(p, 0, 0)
    _run(p, "", size=1)
    return table


def add_section_heading(doc, text, size=14, color=NAVY_DEEP, space_before=12):
    p = doc.add_paragraph()
    _tight(p, space_before, 2)
    _run(p, text, size=size, bold=True, color=color)
    pPr = p._p.get_or_add_pPr()
    pPr.append(parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'<w:bottom w:val="single" w:sz="8" w:space="4" w:color="{_c(BLUE_ACCENT)}"/>'
        f"</w:pBdr>"
    ))
    return p


def add_callout(doc, text, bg=BG_LIGHT, bar_color=NAVY_DEEP, text_color=TEXT_BODY, size=9.5, bold=False):
    table = doc.add_table(rows=1, cols=1)
    _set_col_widths(table, [CONTENT_WIDTH])
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, bg)
    set_cell_margins(cell, top=110, bottom=110, left=200, right=160)
    set_cell_borders(cell, left={"sz": 24, "val": "single", "color": bar_color})
    p = cell.paragraphs[0]
    _tight(p, 0, 0)
    _run(p, text, size=size, color=text_color, bold=bold)
    return table


def add_kpi_cards(doc, cards):
    """cards: [(label, value_text), ...] 4개."""
    n = len(cards)
    table = doc.add_table(rows=1, cols=n)
    _set_col_widths(table, [Inches(CONTENT_WIDTH.inches / n)] * n)
    for i, (label, value) in enumerate(cards):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, BG_LIGHT)
        set_cell_margins(cell, top=140, bottom=140, left=100, right=100)
        set_cell_borders(
            cell,
            top={"sz": 20, "val": "single", "color": BLUE_ACCENT},
            left={"sz": 4, "val": "single", "color": GRAY_BORDER},
            right={"sz": 4, "val": "single", "color": GRAY_BORDER},
            bottom={"sz": 4, "val": "single", "color": GRAY_BORDER},
        )
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p1 = cell.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _tight(p1, 0, 1)
        _run(p1, value, size=17, bold=True, color=NAVY_PRIMARY)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _tight(p2, 0, 0)
        _run(p2, label, size=8.3, color=TEXT_SUB)
    return table


def add_data_table(doc, headers, rows, col_widths=None, right_align_cols=()):
    ncols = len(headers)
    if col_widths is None:
        col_widths = [Inches(CONTENT_WIDTH.inches / ncols)] * ncols
    table = doc.add_table(rows=1 + len(rows), cols=ncols)
    _set_col_widths(table, col_widths)

    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        set_cell_shading(cell, NAVY_PRIMARY)
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        set_cell_borders(
            cell,
            top={"sz": 4, "val": "single", "color": NAVY_PRIMARY},
            bottom={"sz": 4, "val": "single", "color": NAVY_PRIMARY},
            left={"sz": 4, "val": "single", "color": NAVY_PRIMARY},
            right={"sz": 4, "val": "single", "color": NAVY_PRIMARY},
        )
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _tight(p, 2, 2)
        _run(p, h, size=9.3, bold=True, color=WHITE)

    for i, row in enumerate(rows):
        bg = BG_LIGHT2 if i % 2 == 0 else WHITE
        for j, val in enumerate(row):
            cell = table.rows[i + 1].cells[j]
            set_cell_shading(cell, bg)
            set_cell_margins(cell, top=65, bottom=65, left=100, right=100)
            set_cell_borders(cell, bottom={"sz": 4, "val": "single", "color": GRAY_BORDER})
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if j in right_align_cols else WD_ALIGN_PARAGRAPH.LEFT
            _tight(p, 1, 1)
            _run(p, str(val), size=9.3)
    return table


def add_risk_table(doc, rows):
    """rows: [{"label", "value_text", "severity" ("상"/"중"/"하"/None), "desc"}]"""
    headers = ["항목", "값", "심각도", "설명"]
    col_widths = [Inches(1.25), Inches(1.05), Inches(0.85), Inches(3.55)]
    table = doc.add_table(rows=1 + len(rows), cols=4)
    _set_col_widths(table, col_widths)

    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        set_cell_shading(cell, NAVY_PRIMARY)
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _tight(p, 2, 2)
        _run(p, h, size=9.3, bold=True, color=WHITE)

    for i, row in enumerate(rows):
        sev = row.get("severity")
        sev_bg, sev_text = SEVERITY_COLOR.get(sev, (BG_LIGHT2, TEXT_SUB))
        base_bg = WHITE if i % 2 == 0 else BG_LIGHT
        values = [row["label"], row["value_text"], sev or "확인되지 않음", row["desc"]]
        for j, val in enumerate(values):
            cell = table.rows[i + 1].cells[j]
            set_cell_shading(cell, sev_bg if j == 2 else base_bg)
            set_cell_margins(cell, top=65, bottom=65, left=100, right=100)
            set_cell_borders(cell, bottom={"sz": 4, "val": "single", "color": GRAY_BORDER})
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in (1, 2) else WD_ALIGN_PARAGRAPH.LEFT
            _tight(p, 1, 1)
            _run(p, str(val), size=9.3, bold=(j == 2), color=(sev_text if j == 2 else TEXT_BODY))
    return table


# ─────────────────────────────────────────────────────────────────────────
# 리스크 심각도 규칙 (팀 정책에 맞춰 조정 가능한 기본 임계값 — 값이 없으면
# "확인되지 않음"만 반환하고 임의로 등급을 지어내지 않는다)
# ─────────────────────────────────────────────────────────────────────────
def _classify_tariff(rate):
    if rate is None:
        return None
    return "하" if rate <= 5 else ("중" if rate <= 10 else "상")


def _classify_concentration(pct):
    if pct is None:
        return None
    return "하" if pct < 50 else ("중" if pct < 70 else "상")


def _tariff_desc(rate, sev):
    if rate is None:
        return "관세율 정보가 확인되지 않아 별도 확인이 필요합니다."
    if sev == "하":
        return f"관세율이 {fmt(rate, '%')}로 낮은 편이라 충분히 감당 가능한 수준입니다."
    if sev == "중":
        return f"관세율이 {fmt(rate, '%')}로 보통 수준이며, FTA 협정세율 적용 여부를 확인하면 부담을 낮출 수 있습니다."
    return f"관세율이 {fmt(rate, '%')}로 다소 높지만, 원가 구조에 따라 감당 가능한 수준인지 검토가 필요합니다."


def _concentration_desc(pct, sev):
    if pct is None:
        return "상위 3개국 점유율 정보가 확인되지 않아 별도 확인이 필요합니다."
    if sev == "하":
        return f"상위 3개국 점유율이 {fmt(pct, '%')}로 낮아, 소수 국가의 독점도가 크지 않은 감당 가능한 구조입니다."
    if sev == "중":
        return f"상위 3개국 점유율이 {fmt(pct, '%')}로 보통 수준이며, 경쟁 강도를 감안한 전략이 필요합니다."
    return f"상위 3개국 점유율이 {fmt(pct, '%')}로 높은 편이지만, 틈새 포지셔닝에 따라 감당 가능한지 검토가 필요합니다."


# ─────────────────────────────────────────────────────────────────────────
# 문서 준비
# ─────────────────────────────────────────────────────────────────────────
def _setup_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(_c(TEXT_BODY))
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.paragraph_format.space_before = Pt(2)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.12


# ─────────────────────────────────────────────────────────────────────────
# 표지
# ─────────────────────────────────────────────────────────────────────────
def _build_cover_page(doc: Document, data: dict) -> None:
    # 상단 투톤 액센트 배너
    banner = doc.add_table(rows=1, cols=2)
    _set_col_widths(banner, [Inches(CONTENT_WIDTH.inches * 0.6), Inches(CONTENT_WIDTH.inches * 0.4)])
    left_cell, right_cell = banner.rows[0].cells
    set_cell_shading(left_cell, NAVY_PRIMARY)
    set_cell_shading(right_cell, BLUE_ACCENT)
    for c in (left_cell, right_cell):
        set_cell_margins(c, top=6, bottom=6, left=0, right=0)
        _run(c.paragraphs[0], "", size=1)

    doc.add_paragraph().paragraph_format.space_after = Pt(20)

    badge_p = doc.add_paragraph()
    _tight(badge_p, 0, 4)
    _run(badge_p, "AI TRADE INTELLIGENCE REPORT | 무역 기회 발굴 분석", size=9.5, bold=True, color=BLUE_ACCENT)

    title_p1 = doc.add_paragraph()
    _tight(title_p1, 16, 0)
    _run(title_p1, "Blue Ocean Finder", size=38, bold=True, color=NAVY_DEEP)

    title_p2 = doc.add_paragraph()
    _tight(title_p2, 0, 10)
    _run(title_p2, "시장 보고서", size=38, bold=True, color=NAVY_PRIMARY)

    add_divider(doc, width=Inches(2.5), color=BLUE_ACCENT)

    subtitle_p = doc.add_paragraph()
    _tight(subtitle_p, 10, 24)
    _run(subtitle_p, "무역 통계 데이터 기반 글로벌 틈새 유망 시장 분석 및 의사결정 브리프", size=11, color=TEXT_SUB)

    # 메타데이터 카드
    item_label = safe(data.get("item_name"))
    hs_code = data.get("hs_code")
    if data.get("item_name") and hs_code:
        item_label = f"{data['item_name']} (HS {hs_code})"

    rows_meta = [
        ("수출 품목", item_label),
        ("추천 국가", safe(data.get("country"))),
        ("분석 엔진", "Predator Engine"),
        ("작성 프로젝트", safe(data.get("project_name"))),
        ("조사 기준일", safe(data.get("date"))),
        ("발행 일자", date.today().isoformat()),
    ]
    meta = doc.add_table(rows=len(rows_meta), cols=2)
    _set_col_widths(meta, [Inches(1.9), Inches(CONTENT_WIDTH.inches - 1.9)])
    for i, (label, value) in enumerate(rows_meta):
        bg = BG_LIGHT if i % 2 == 0 else BG_LIGHT2
        label_cell, value_cell = meta.rows[i].cells
        for cell in (label_cell, value_cell):
            set_cell_shading(cell, bg)
            set_cell_margins(cell, top=90, bottom=90, left=160, right=120)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        # 좌측 전체에 이어지는 네이비 세로 액센트 바
        set_cell_borders(label_cell, left={"sz": 24, "val": "single", "color": NAVY_DEEP})
        lp = label_cell.paragraphs[0]
        _tight(lp, 0, 0)
        _run(lp, label, size=9.5, bold=True, color=NAVY_PRIMARY)
        vp = value_cell.paragraphs[0]
        _tight(vp, 0, 0)
        _run(vp, str(value), size=9.5, color=TEXT_BODY)

    footer_p = doc.add_paragraph()
    _tight(footer_p, 22, 0)
    team_label = safe(data.get("project_name"), na="Blue Ocean Finder")
    _run(footer_p, f"{team_label} · PROJECT TEAM", size=8.3, color=TEXT_SUB)

    doc.add_page_break()


# ─────────────────────────────────────────────────────────────────────────
# 1. 요약
# ─────────────────────────────────────────────────────────────────────────
def _build_section_summary(doc: Document, data: dict) -> None:
    add_section_heading(doc, "1. 요약")

    rank_val = data.get("rank")  # 정식 스키마 외 선택 필드 — 없으면 항상 "확인되지 않음"
    rank_text = f"{int(rank_val)}위" if isinstance(rank_val, (int, float)) else "확인되지 않음"

    add_kpi_cards(
        doc,
        [
            ("Blue Ocean Score", fmt(data.get("blue_ocean_score"), "점")),
            ("한국 점유율", fmt(data.get("korea_market_share"), "%")),
            ("최근 1년 성장률", fmt(data.get("growth_1y"), "%", signed=True)),
            ("종합 순위", rank_text),
        ],
    )

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    country = safe(data.get("country"))
    item_name = safe(data.get("item_name"))
    score_txt = fmt(data.get("blue_ocean_score"), "점")
    growth_txt = fmt(data.get("growth_1y"), "%", signed=True)
    share_txt = fmt(data.get("korea_market_share"), "%")
    conclusion = (
        f"{country}은(는) {item_name} 품목 기준 Blue Ocean Score {score_txt}으로 분석된 시장입니다. "
        f"최근 1년 수입 성장률은 {growth_txt}이며, 한국 점유율은 {share_txt}로 나타나 "
        f"추가 진입 여지가 있는지 검토할 만한 것으로 판단됩니다."
    )
    add_callout(doc, conclusion, bg=BG_LIGHT, bar_color=NAVY_DEEP, size=9.7)


# ─────────────────────────────────────────────────────────────────────────
# 2. 이 시장을 선택한 이유
# ─────────────────────────────────────────────────────────────────────────
def _build_section_reason(doc: Document, data: dict) -> None:
    add_section_heading(doc, "2. 이 시장을 선택한 이유")

    growth_1y = data.get("growth_1y")
    export_gap = data.get("export_gap")
    top3 = data.get("top3_concentration")

    reasons = [
        (
            "최근 1년 수입 성장률",
            f"{fmt(growth_1y, '%', signed=True)} — "
            + ("이 시장은 최근 확대되는 추세로 보입니다." if isinstance(growth_1y, (int, float)) and growth_1y > 0
               else "최근 성장률 자료를 확인한 뒤 진입 시점을 판단하는 것이 좋습니다."),
        ),
        (
            "수출 격차 (Export Gap)",
            f"{fmt(export_gap, '%p', signed=True)} — "
            "한국의 세계 점유율 대비 이 시장에서의 점유율이 낮아, 상대적으로 진입 여지가 있는 것으로 해석됩니다.",
        ),
        (
            "상위 3개국 점유율",
            f"{fmt(top3, '%')} — "
            "소수 국가의 독점도를 가늠하는 지표로, 값이 낮을수록 신규 진입 여지가 큽니다.",
        ),
    ]
    for label, desc in reasons:
        p = doc.add_paragraph()
        _tight(p, 3, 3)
        _run(p, f"● {label}: ", size=9.6, bold=True, color=NAVY_PRIMARY)
        _run(p, desc, size=9.6, color=TEXT_BODY)

    appendix_heading = doc.add_paragraph("[부록 A] 산출 배경 보기", style="Heading 3")
    for r in appendix_heading.runs:
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor.from_string(_c(NAVY_PRIMARY))
        r.font.bold = True
        _apply_font(r)
    _tight(appendix_heading, 8, 2)

    bg_text = (
        f"Market Opportunity Score: {fmt(data.get('market_opportunity_score'), '점')}   ·   "
        f"Penetration Opportunity Score: {fmt(data.get('penetration_opportunity_score'), '점')}   ·   "
        f"Blue Ocean Score: {fmt(data.get('blue_ocean_score'), '점')}   ·   "
        f"3년 CAGR: {fmt(data.get('cagr_3y'), '%', signed=True)}"
    )
    add_callout(doc, bg_text, bg=BG_LIGHT2, bar_color=BLUE_ACCENT, text_color=TEXT_SUB, size=9)


# ─────────────────────────────────────────────────────────────────────────
# 3. 시장 현황 스냅샷
# ─────────────────────────────────────────────────────────────────────────
def _build_section_snapshot(doc: Document, data: dict) -> None:
    add_section_heading(doc, "3. 시장 현황 스냅샷")

    rows = [
        ("최근 1년 수입 성장률", fmt(data.get("growth_1y"), "%", signed=True)),
        ("최근 3년 CAGR", fmt(data.get("cagr_3y"), "%", signed=True)),
        ("한국 점유율 (이 시장)", fmt(data.get("korea_market_share"), "%")),
        ("한국 점유율 (글로벌)", fmt(data.get("global_korea_share"), "%")),
        ("수출 격차 (Export Gap)", fmt(data.get("export_gap"), "%p", signed=True)),
    ]
    add_data_table(
        doc,
        headers=["지표", "값"],
        rows=rows,
        col_widths=[Inches(4.2), Inches(2.5)],
        right_align_cols=(1,),
    )
    caption = doc.add_paragraph()
    _tight(caption, 4, 6)
    _run(caption, f"자료: 무역 통계 기반 산출값 (조사 기준일: {safe(data.get('date'))})", size=8, color=TEXT_SUB, italic=True)


# ─────────────────────────────────────────────────────────────────────────
# 4. 리스크 및 진입장벽
# ─────────────────────────────────────────────────────────────────────────
def _build_section_risk(doc: Document, data: dict) -> None:
    add_section_heading(doc, "4. 리스크 및 진입장벽")

    tariff = data.get("tariff_rate")
    top3 = data.get("top3_concentration")
    tariff_sev = _classify_tariff(tariff)
    top3_sev = _classify_concentration(top3)

    rows = [
        {
            "label": "관세율",
            "value_text": fmt(tariff, "%"),
            "severity": tariff_sev,
            "desc": _tariff_desc(tariff, tariff_sev),
        },
        {
            "label": "경쟁 강도\n(상위 3개국 점유율)",
            "value_text": fmt(top3, "%"),
            "severity": top3_sev,
            "desc": _concentration_desc(top3, top3_sev),
        },
        {
            "label": "비관세 규제",
            "value_text": "확인되지 않음",
            "severity": None,
            "desc": "비관세 규제 세부 항목은 이번 산출 범위에 포함되지 않아 별도 확인이 필요합니다.",
        },
    ]
    add_risk_table(doc, rows)


# ─────────────────────────────────────────────────────────────────────────
# 5. 데이터 한계
# ─────────────────────────────────────────────────────────────────────────
def _build_section_limitation(doc: Document, data: dict) -> None:
    add_section_heading(doc, "5. 데이터 한계")
    text = (
        f"본 리포트의 관세율 수치는 {safe(data.get('date'))} 기준 조사된 값이며, "
        "결제 관행·정부지원사업 등 일부 보조 지표는 이번 산출 범위에 포함되지 않았습니다."
    )
    add_callout(doc, text, bg=BG_LIGHT2, bar_color=TEXT_SUB, text_color=TEXT_SUB, size=9)


# ─────────────────────────────────────────────────────────────────────────
# 6. 다음 액션 체크리스트
# ─────────────────────────────────────────────────────────────────────────
def _build_section_next_actions(doc: Document, data: dict) -> None:
    add_section_heading(doc, "6. 다음 액션 체크리스트")

    country = safe(data.get("country"))
    item_name = safe(data.get("item_name"))
    tariff_txt = fmt(data.get("tariff_rate"), "%")

    rows = [
        (
            "1단계",
            "데이터 재검증",
            f"{item_name}의 {country} 수출 통계와 관세율({tariff_txt})을 최신 자료로 다시 확인합니다.",
        ),
        (
            "2단계",
            "현지 유통·규제 조사",
            f"{country}의 유통 구조, 인증 요건, 비관세 규제를 조사합니다.",
        ),
        (
            "3단계",
            "소량 샘플 테스트",
            "소량 샘플 오더로 현지 반응과 통관 절차를 직접 확인합니다.",
        ),
    ]
    add_data_table(
        doc,
        headers=["단계", "실행 항목", "확인 포인트"],
        rows=rows,
        col_widths=[Inches(0.8), Inches(1.7), Inches(4.2)],
    )


# ─────────────────────────────────────────────────────────────────────────
# 7. 접촉 채널 (참고용)
# ─────────────────────────────────────────────────────────────────────────
def _build_section_contacts(doc: Document, data: dict) -> None:
    add_section_heading(doc, "7. 접촉 채널")

    badge_p = doc.add_paragraph()
    _tight(badge_p, 0, 4)
    _run(badge_p, "참고용", size=8.5, bold=True, color=BLUE_ACCENT)

    country = safe(data.get("country"))
    item_name = safe(data.get("item_name"))

    rows = [
        ("KOTRA 무역관", f"{country} 지역을 담당하는 KOTRA 해외 무역관에 문의해 최신 시장 정보와 바이어 연결을 요청할 수 있습니다."),
        ("협회 · 전시회", f"{item_name} 관련 국내외 산업 협회 및 전시회 정보를 확인해 참가를 검토할 수 있습니다."),
        ("온라인 B2B 플랫폼", "Alibaba 등 온라인 B2B 플랫폼과 현지 조달 플랫폼에서 바이어 탐색을 시작할 수 있습니다."),
    ]
    add_data_table(
        doc,
        headers=["채널", "안내"],
        rows=rows,
        col_widths=[Inches(1.6), Inches(5.1)],
    )
    note = doc.add_paragraph()
    _tight(note, 4, 2)
    _run(note, "※ 위 채널은 일반적인 참고 가이드이며, 실제 담당 기관·전시회명은 별도 확인이 필요합니다.", size=8, color=TEXT_SUB, italic=True)


# ─────────────────────────────────────────────────────────────────────────
# 8. 종합 결론 및 방향성 (참고용)
# ─────────────────────────────────────────────────────────────────────────
def _build_section_conclusion(doc: Document, data: dict) -> None:
    add_section_heading(doc, "8. 종합 결론 및 방향성")

    badge_p = doc.add_paragraph()
    _tight(badge_p, 0, 4)
    _run(badge_p, "참고용", size=8.5, bold=True, color=BLUE_ACCENT)

    score = data.get("blue_ocean_score")
    if score is None:
        grade, bg, text_color = "확인되지 않음", BG_LIGHT2, TEXT_SUB
        grade_desc = "Blue Ocean Score 값이 없어 등급을 판정할 수 없습니다."
    elif score >= 80:
        grade, bg, text_color = "적극 검토 가능", STATUS_GOOD_BG, STATUS_GOOD_TEXT
        grade_desc = f"Blue Ocean Score {fmt(score, '점')}로 80점 이상 구간에 해당해 적극 검토가 가능한 수준입니다."
    elif score >= 60:
        grade, bg, text_color = "신중 접근", STATUS_WARN_BG, STATUS_WARN_TEXT
        grade_desc = f"Blue Ocean Score {fmt(score, '점')}로 60~79점 구간에 해당해 신중한 접근이 필요합니다."
    else:
        grade, bg, text_color = "보류", STATUS_RISK_BG, STATUS_RISK_TEXT
        grade_desc = f"Blue Ocean Score {fmt(score, '점')}로 60점 미만 구간에 해당해 진출을 보류하는 것을 검토할 만합니다."

    p = doc.add_paragraph()
    _tight(p, 4, 2)
    _run(p, f"판정: {grade}", size=11, bold=True, color=text_color)
    add_callout(doc, grade_desc, bg=bg, bar_color=text_color, text_color=text_color, size=9.7)

    disclaimer = doc.add_paragraph()
    _tight(disclaimer, 6, 2)
    _run(
        disclaimer,
        "※ 본 판정은 산출 데이터를 기준으로 한 참고용 결과이며, 최종 결정은 팀 내 추가 검토가 필요합니다.",
        size=8.5,
        color=TEXT_SUB,
        italic=True,
    )


# ─────────────────────────────────────────────────────────────────────────
# A. 산출 근거
# ─────────────────────────────────────────────────────────────────────────
def _build_appendix_formula(doc: Document) -> None:
    add_section_heading(doc, "A. 산출 근거")

    note = doc.add_paragraph()
    _tight(note, 2, 4)
    _run(note, "본 표는 산출 방법론을 설명하기 위한 표준 공식이며, 세부 가중치·정규화 방식은 프로젝트 계산 로직을 따릅니다.", size=8.3, color=TEXT_SUB, italic=True)

    rows = [
        ("Market Opportunity Score", "수요 측 지표 정규화 가중합 (0~100)", "시장 규모·성장률·트렌드·환율 등 수요 측 지표를 종합한 점수"),
        ("Penetration Opportunity Score", "공급 여지 지표 정규화 가중합 (0~100)", "한국 점유율·경쟁 집중도·관세 등 공급 여지 지표를 종합한 점수"),
        ("Blue Ocean Score", "√(Market Opportunity Score × Penetration Opportunity Score)", "두 점수의 기하평균 — 한쪽이 매우 낮으면 전체 점수도 함께 낮아짐"),
        ("Export Gap", "Global Korea Share(%) − Korea Market Share(%)", "한국의 세계 점유율 대비 이 시장에서의 점유율 격차(%p)"),
        ("3년 CAGR", "((T년 값 ÷ T-3년 값)^(1/3) − 1) × 100", "최근 3개년 연평균 성장률(%)"),
    ]
    add_data_table(
        doc,
        headers=["지표", "산출식", "설명"],
        rows=rows,
        col_widths=[Inches(1.7), Inches(2.6), Inches(2.4)],
    )


# ─────────────────────────────────────────────────────────────────────────
# 공개 함수
# ─────────────────────────────────────────────────────────────────────────
def generate_blue_ocean_report(data: dict, filename: str = "Blue_Ocean_Finder_Report.docx") -> str:
    """`data`만으로 Blue Ocean Finder Word 보고서를 생성해 `filename`에 저장한다."""
    doc = Document()
    _setup_document(doc)

    _build_cover_page(doc, data)  # 내부에서 add_page_break() 호출
    _build_section_summary(doc, data)
    _build_section_reason(doc, data)
    _build_section_snapshot(doc, data)
    _build_section_risk(doc, data)
    _build_section_limitation(doc, data)
    _build_section_next_actions(doc, data)
    _build_section_contacts(doc, data)
    _build_section_conclusion(doc, data)
    _build_appendix_formula(doc)

    doc.save(filename)
    return filename


if __name__ == "__main__":
    # 실행 예시 — 테스트용 데이터. 실제 사용 시 계산 파이프라인이 만든 딕셔너리로 교체할 것.
    sample_data = {
        "country": "멕시코",
        "blue_ocean_score": 82.4,
        "market_opportunity_score": 78.1,
        "penetration_opportunity_score": 87.0,
        "growth_1y": 12.3,
        "cagr_3y": 9.8,
        "korea_market_share": 1.2,
        "global_korea_share": 16.6,
        "export_gap": 15.4,
        "tariff_rate": 0.0,
        "top3_concentration": 58.3,
        "item_name": "태양광 모듈",
        "hs_code": "854140",
        "project_name": "Blue Ocean Finder",
        "date": "2026-09-20",
    }
    out = generate_blue_ocean_report(sample_data)
    print(f"보고서 생성 완료: {out}")
